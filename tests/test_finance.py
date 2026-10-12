import io
from datetime import date, datetime
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from brandwise_system.db import DB
from brandwise_system.finance import Filters, Finance, to_grosze, zl
from brandwise_system.legacy_import import ImportErrorPL, check_totals, import_legacy, parse_dump

from .conftest import ORIGIN, login

H = {"Origin": ORIGIN}


@pytest.fixture
def fin(tmp_path):
    return Finance(DB(tmp_path / "f.db"))


def test_money_helpers():
    assert to_grosze("1 234,50") == 123450 and to_grosze("99.9") == 9990 and to_grosze("") == 0
    assert zl(123450) == "1 234,50 zł" and zl(-5) == "−0,05 zł" and zl(0) == "0,00 zł"


def test_orders_filters_and_summary(fin):
    a, b = fin.add_client("Allianz"), fin.add_client("LUX MED")
    base = dict(hashtag="#16", status="Projektowanie", principal="Anna", qty=1, costs=[])
    fin.save_order({**base, "client_id": a, "payer_id": a, "day": "2026-10-03", "title": "Kartka", "price": 100000,
                    "invoice_status": "Brak", "costs": [["druk", 20000]]})
    fin.save_order({**base, "client_id": b, "payer_id": a, "day": "2026-10-10", "title": "Baner", "price": 5000,
                    "qty": 3, "invoice_status": "Faktura"})
    fin.save_order({**base, "client_id": b, "payer_id": b, "day": "2026-09-20", "title": "Stara", "price": 10000,
                    "invoice_status": "Faktura"})
    t = fin.totals(Filters(day_from="2026-10-01", day_to="2026-10-31"))
    assert t == {"n": 2, "total": 115000, "costs": 20000, "net": 95000}
    assert fin.totals(Filters(client_id=(b,)))["n"] == 2
    assert fin.totals(Filters(payer_id=(a,)))["n"] == 2
    assert fin.totals(Filters(q="ban"))["n"] == 1
    fin.add_cost("2026-10", "Biuro", 30000)
    s = fin.month_summary("2026-10")
    assert s["revenue"] == 95000 and s["fixed_costs"] == 30000 and s["profit"] == 65000
    assert s["invoiced_pct_count"] == 50 and s["to_invoice_count"] == 1 and s["to_invoice_total"] == 100000
    assert fin.by_client("2026-10") == [("Allianz", 80000), ("LUX MED", 15000)]
    assert fin.year_table()["2026"][8] == 10000
    assert fin.copy_costs("2026-10", "2026-11") == 1


def test_validation(fin):
    with pytest.raises(ValueError):
        fin.save_order({"title": "", "client_id": 1, "payer_id": 1, "day": "2026-10-01"})
    c = fin.add_client("X")
    with pytest.raises(ValueError):
        fin.save_order({"title": "t", "client_id": c, "payer_id": c, "day": "32-01-2026"})


def test_report_expiry_and_revoke(fin):
    tok = fin.create_report(Filters(hashtag="#1"), "Raport", "2000-01-01")
    assert fin.report(tok) is None                              # wygasł
    tok = fin.create_report(Filters(hashtag="#1"), "Raport")
    assert fin.report(tok) is not None
    fin.revoke_report(tok)
    assert fin.report(tok) is None


def test_routes_full_flow(make_client):
    c, app = make_client()
    F = app.state.finance
    assert c.get("/zlecenia").status_code == 303                # bez logowania
    login(c, app)
    assert c.post("/klienci", data={"name": "PGNiG"}, headers=H).status_code == 303
    cid = F.clients()[0]["id"]
    r = c.post("/zlecenia", headers=H, data={
        "day": date.today().isoformat(), "hashtag": "#16", "title": "Ulotka", "client_id": cid, "payer_id": "",
        "status": "Projektowanie", "invoice_status": "Brak", "price": "1 200,00", "qty": "2",
        "koszt_nazwa": ["druk", ""], "koszt_kwota": ["300,50", ""], "wroc": "/zlecenia?klient=1"})
    assert r.status_code == 303 and r.headers["location"].startswith("/zlecenia?klient=1")
    o = F.orders(Filters())[0]
    assert (o["total"], o["costs_total"], o["payer_id"]) == (240000, 30050, cid)
    html = c.get("/zlecenia").text
    assert "Ulotka" in html and "2 400,00 zł" in html
    assert "Ulotka" in c.get(f"/zlecenia/{o['id']}").text
    assert "Kopia zlecenia" in c.get(f"/zlecenia/{o['id']}/duplikuj").text
    # zbiorcza zmiana
    assert c.post("/zlecenia/zbiorczo", headers=H, data={"id": [o["id"]], "invoice_status": "Faktura"}).status_code == 303
    assert F.order(o["id"])["invoice_status"] == "Faktura"
    # przekierowanie tylko wewnątrz systemu
    r = c.post(f"/zlecenia/{o['id']}/usun", headers=H, data={"wroc": "https://zly.example/"})
    assert r.headers["location"].startswith("/zlecenia")
    assert F.order(o["id"]) is None
    for p in ("/finanse", "/koszty", "/klienci", "/ustawienia", "/raporty", "/"):
        assert c.get(p).status_code == 200, p


def test_public_report_and_xlsx(make_client):
    c, app = make_client()
    F = app.state.finance
    cid = F.add_client("Allianz")
    F.save_order({"day": "2026-10-01", "title": "Kartka świąteczna", "client_id": cid, "payer_id": cid,
                  "price": 50000, "qty": 1, "status": "Zakończone", "invoice_status": "Faktura",
                  "costs": [["tajny koszt", 12300]]})
    login(c, app)
    r = c.post("/raporty", headers=H, data={"filters": '{"klient": ["%d"]}' % cid, "title": "Allianz – październik"})
    token = r.headers["location"].split("nowy=")[1]
    c.cookies.clear()                                            # klient bez logowania
    html = c.get(f"/r/{token}").text
    assert "Kartka świąteczna" in html and "500,00 zł" in html
    assert "tajny koszt" not in html and "123,00" not in html and "Faktura" not in html
    x = c.get(f"/r/{token}/xlsx")
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["C2"].value == "Kartka świąteczna" and ws["I2"].value == 500.0
    assert c.get("/r/zly-token").status_code == 404
    assert c.get("/raporty").status_code == 303                  # lista linków tylko po zalogowaniu


EOL = "\r\n"


def legacy_dump() -> str:
    ts = int(datetime(2026, 9, 30, tzinfo=ZoneInfo("Europe/Warsaw")).timestamp())   # północ czasu polskiego
    return EOL.join([
        "DROP TABLE IF EXISTS `clients`;",
        "CREATE TABLE `clients` (\n  `id` int unsigned NOT NULL AUTO_INCREMENT,\n  `name` varchar(255) NOT NULL,\n"
        "  `created_at` timestamp NULL DEFAULT NULL,\n  `updated_at` timestamp NULL DEFAULT NULL,\n"
        "  `deleted_at` timestamp NULL DEFAULT NULL,\n  PRIMARY KEY (`id`)\n) ENGINE=InnoDB;",
        "INSERT INTO `clients` VALUES ('1','PGNiG Obrót Detaliczny',NULL,NULL,NULL),\n"
        "('2','Allianz',NULL,NULL,'2020-01-01 00:00:00'); ",
        "CREATE TABLE `dates` (\n  `id` int unsigned NOT NULL,\n  `value` varchar(255) NOT NULL,\n"
        "  `created_at` timestamp NULL,\n  `updated_at` timestamp NULL,\n  PRIMARY KEY (`id`)\n) ENGINE=InnoDB;",
        "INSERT INTO `dates` VALUES ('5','09/2026',NULL,NULL); ",
        "CREATE TABLE `costs` (\n  `id` int unsigned NOT NULL,\n  `name` varchar(255) NOT NULL,\n"
        "  `value` varchar(255) NOT NULL,\n  `date_id` int unsigned NOT NULL,\n  `created_at` timestamp NULL,\n"
        "  `updated_at` timestamp NULL,\n  PRIMARY KEY (`id`)\n) ENGINE=InnoDB;",
        "INSERT INTO `costs` VALUES ('1','Biuro','2500.5','5',NULL,NULL),\n('2','Sierota','10','99',NULL,NULL); ",
        "CREATE TABLE `orders` (\n  `id` int unsigned NOT NULL,\n  `hashtag` varchar(255) NOT NULL,\n"
        "  `client` varchar(255) NOT NULL,\n  `invoice` varchar(255) NOT NULL,\n  `data` varchar(255) NOT NULL,\n"
        "  `title` varchar(255) NOT NULL,\n  `status` varchar(255) NOT NULL,\n  `principal` varchar(255) NOT NULL,\n"
        "  `account` varchar(255) NOT NULL,\n  `price` varchar(255) NOT NULL,\n  `pcs` varchar(255) NOT NULL,\n"
        "  `sum` varchar(255) NOT NULL,\n  `costs_data` varchar(1024) DEFAULT NULL,\n"
        "  `costs_sum` varchar(255) DEFAULT NULL,\n  `invoice_status` varchar(255) NOT NULL,\n"
        "  `created_by` varchar(255) NOT NULL,\n  `created_at` timestamp NULL,\n  `updated_at` timestamp NULL,\n"
        "  `deleted_at` timestamp NULL,\n  PRIMARY KEY (`id`)\n) ENGINE=InnoDB;",
        f"INSERT INTO `orders` VALUES ('10','#16 - umowa','1','1','{ts}','Ulotka \\'jesień\\'\\r\\n2026','Zakończone',"
        f"'Paweł','3','1500.5','2','3001','{{\\\"druk\\\":\\\"300\\\"}}','300','Faktura','3','2026-09-30 10:00:00',NULL,NULL),\n"
        f"('11','0','1','2','01/10/2026','Baner','Wycena','','4','100','1','100',NULL,NULL,'Brak','4',NULL,NULL,NULL),\n"
        f"('12','x','7','7','{ts}','Bez klienta','Wycena','','3','0','1','0',NULL,NULL,'Brak','3',NULL,NULL,NULL); ",
        "CREATE TABLE `settings` (\n  `id` int unsigned NOT NULL,\n  `name` varchar(255) NOT NULL,\n"
        "  `value` varchar(255) NOT NULL,\n  `created_at` timestamp NULL,\n  `updated_at` timestamp NULL,\n"
        "  PRIMARY KEY (`id`)\n) ENGINE=InnoDB;",
        "INSERT INTO `settings` VALUES ('1','status','[\\\"Wycena\\\",\\\"Zako\\\\u0144czone\\\"]',NULL,NULL),\n"
        "('2','invoice_status','[\\\"Brak\\\",\\\"Faktura\\\"]',NULL,NULL); ",
        "CREATE TABLE `users` (\n  `id` int unsigned NOT NULL,\n  `name` varchar(255) NOT NULL,\n"
        "  `email` varchar(255) NOT NULL,\n  `password` varchar(255) NOT NULL,\n  `remember_token` varchar(100) NULL,\n"
        "  `created_at` timestamp NULL,\n  `updated_at` timestamp NULL,\n  PRIMARY KEY (`id`)\n) ENGINE=InnoDB;",
        "INSERT INTO `users` VALUES ('3','Łukasz','lukasz@brandwise.pl','$2y$10$x',NULL,NULL,NULL),\n"
        "('4','Ola','ola@brandwise.pl','$2y$10$y',NULL,NULL,NULL); ",
        "",
    ])


def test_legacy_import(fin):
    data = parse_dump(legacy_dump())
    assert len(data["orders"]) == 3
    assert data["orders"][0]["title"] == "Ulotka 'jesień'\r\n2026"
    stats = import_legacy(fin, data)
    assert stats["klienci"] == 2 and stats["zlecenia"] == 3 and stats["koszty"] == 1
    assert any("klient 7" in p for p in stats["pominiete"]) and any("koszt 2" in p for p in stats["pominiete"])
    o = fin.order(10)
    assert o["day"] == "2026-09-30"                             # północ w Polsce, nie dzień wcześniej (UTC)
    assert (o["price"], o["qty"], o["total"], o["costs_total"]) == (150050, 2.0, 300100, 30000)
    assert o["account_legacy"] == "Łukasz" and o["hashtag"] == "#16 - umowa"
    o = fin.order(11)
    assert o["day"] == "2026-10-01" and o["hashtag"] == "" and o["payer_id"] == 2 and o["account_legacy"] == "Ola"
    assert fin.clients(with_deleted=False)[0]["name"] == "PGNiG Obrót Detaliczny"   # Allianz usunięty
    assert fin.costs("2026-09")[0]["value"] == 250050
    assert fin.get_list("status") == ["Wycena", "Zakończone"]
    assert check_totals(fin) == [("2026", 3, 310100)]
    with pytest.raises(ImportErrorPL):
        import_legacy(fin, data)                                # drugi raz bez --zastap
    import_legacy(fin, data, replace=True)
    assert fin.totals(Filters())["n"] == 3


def test_import_rejects_garbage(fin):
    with pytest.raises(ImportErrorPL):
        parse_dump("SELECT 1;")
