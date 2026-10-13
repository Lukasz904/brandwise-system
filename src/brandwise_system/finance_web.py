"""Widoki panelu finansowego (zlecenia, klienci, koszty, podsumowania, raporty dla klientów)."""

from __future__ import annotations

import io
import json
from datetime import date
from urllib.parse import urlencode

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response

from .finance import Filters, Finance, month_label, prev_month, to_grosze, to_qty, valid_month, zl

router = APIRouter()


def fin(request: Request) -> Finance:
    return request.app.state.finance


def page(request: Request, name: str, **ctx):
    from .app import page_response
    return page_response(request, name, **ctx)


def safe_path(url: str | None, default: str = "/zlecenia") -> str:
    url = url or ""
    if not url.startswith("/") or url.startswith("//") or "\\" in url:
        return default                                  # tylko adresy wewnątrz systemu
    return url


def back(url: str, msg: str = "") -> RedirectResponse:
    url = safe_path(url)
    sep = "&" if "?" in url else "?"
    return RedirectResponse(url + (f"{sep}ok={msg}" if msg else ""), status_code=303)


def this_month() -> str:
    return date.today().strftime("%Y-%m")


def month_last_day(m: str) -> str:
    from .finance import month_bounds
    nxt = date.fromisoformat(month_bounds(m)[1])
    return date.fromordinal(nxt.toordinal() - 1).isoformat()


def month_or_now(m: str) -> str:
    return m if valid_month(m) else this_month()


def client_choices(F: Finance, *keep) -> list:
    """Aktywni klienci + usunięci, do których odwołuje się edytowane zlecenie (żeby nie zgubić płatnika)."""
    active = list(F.clients())
    ids = {c["id"] for c in active}
    extra = [{"id": c["id"], "name": c["name"] + " (usunięty)"} for c in F.clients(with_deleted=True)
             if c["id"] in keep and c["id"] not in ids]
    return active + extra


def filter_choices(F: Finance) -> list:
    return [{"id": c["id"], "name": c["name"] + (" (usunięty)" if c["deleted_at"] else "")}
            for c in F.clients(with_deleted=True)]


def form_order(form) -> dict:
    names, values = form.getlist("koszt_nazwa"), form.getlist("koszt_kwota")
    return {
        "hashtag": form.get("hashtag", ""), "client_id": form.get("client_id") or 0,
        "payer_id": form.get("payer_id") or form.get("client_id") or 0, "day": form.get("day", ""),
        "title": form.get("title", ""), "status": form.get("status", ""), "principal": form.get("principal", ""),
        "price": to_grosze(form.get("price")), "qty": to_qty(form.get("qty")),
        "invoice_status": form.get("invoice_status", ""),
        "costs": [[n, to_grosze(v)] for n, v in zip(names, values)],
        "note": form.get("note", ""),
    }


def log(request: Request, event: str, detail: str = "") -> None:
    from .auth import current_user
    u = current_user(request)
    request.app.state.db.log(event, u["email"] if u else None, detail)


# ---------- zlecenia i prospekty (ta sama struktura, inna „szuflada”) ----------
KINDS = {
    "zlecenie": {"kind": "zlecenie", "base": "/zlecenia", "title": "Zlecenia", "one": "Zlecenie", "new": "Nowe zlecenie",
                 "of": "zleceń"},
    "prospekt": {"kind": "prospekt", "base": "/prospekty", "title": "Prospekty", "one": "Prospekt",
                 "new": "Nowy prospekt", "of": "prospektów"},
}


def _list(request: Request, K: dict):
    f = Filters.from_query(request.query_params)
    f.kind = K["kind"]
    preset = request.query_params.get("okres", "")
    if preset in ("biezacy", "poprzedni"):
        m = this_month() if preset == "biezacy" else prev_month(this_month())
        f.day_from, f.day_to = f"{m}-01", month_last_day(m)
    elif f.empty and "wszystkie" not in request.query_params and K["kind"] == "zlecenie":
        m = this_month()
        f.day_from, f.day_to = f"{m}-01", month_last_day(m)
        preset = "biezacy"
    F = fin(request)
    limit = 2000
    return page(request, "fin_zlecenia.html", K=K, f=f, preset=preset, orders=F.orders(f, limit), totals=F.totals(f),
                limit=limit, clients=filter_choices(F), statuses=F.get_list("status"),
                inv_statuses=F.get_list("invoice_status"), query=urlencode(f.to_query(), doseq=True),
                msg=request.query_params.get("ok", ""))


def order_form(request: Request, K: dict, order=None, duplicate=False):
    F = fin(request)
    data = dict(order) if order else {"day": date.today().isoformat(), "qty": 1, "costs": "[]",
                                      "status": (F.get_list("status") or [""])[0], "invoice_status": "Brak"}
    data["costs_list"] = json.loads(data.get("costs") or "[]")
    if duplicate:
        data["id"] = None
        data["day"] = date.today().isoformat()
    return page(request, "fin_zlecenie.html", K=K, o=data,
                clients=client_choices(F, data.get("client_id"), data.get("payer_id")), statuses=F.get_list("status"),
                inv_statuses=F.get_list("invoice_status"), duplicate=duplicate, error="", today=date.today().isoformat(),
                wroc=safe_path(request.query_params.get("wroc"), K["base"]))


def _get(request: Request, K: dict, oid: int):
    o = fin(request).order(oid)
    if o is None or o["kind"] != K["kind"]:
        raise HTTPException(404, f"Nie ma takiego wpisu w: {K['title']}.")
    return o


async def _save(request: Request, K: dict, oid: int | None):
    form = await request.form()
    F = fin(request)
    try:
        data = form_order(form)
        data["kind"] = K["kind"]
        new_id = F.save_order(data, oid)
    except (ValueError, OverflowError) as e:
        data = {k: form.get(k, "") for k in ("hashtag", "day", "title", "status", "principal", "invoice_status", "note")}
        data.update(id=oid, client_id=int(form.get("client_id") or 0), payer_id=int(form.get("payer_id") or 0),
                    price=0, qty=form.get("qty") or 1, costs_list=[])
        return page(request, "fin_zlecenie.html", K=K, o=data, clients=client_choices(F, data["client_id"], data["payer_id"]),
                    statuses=F.get_list("status"), inv_statuses=F.get_list("invoice_status"), duplicate=False,
                    error=str(e) + " Popraw dane i zapisz ponownie (cena i koszty wymagają ponownego wpisania).",
                    today=date.today().isoformat(), wroc=safe_path(form.get("wroc"), K["base"]))
    log(request, f"{K['kind']}_saved" if oid else f"{K['kind']}_created", str(new_id))
    return back(safe_path(form.get("wroc"), K["base"]), "zapisano")


def register(K: dict) -> None:
    base = K["base"]

    async def list_(request: Request):
        return _list(request, K)

    async def new(request: Request):
        return order_form(request, K)

    async def edit(request: Request, oid: int):
        return order_form(request, K, _get(request, K, oid))

    async def duplicate(request: Request, oid: int):
        return order_form(request, K, _get(request, K, oid), duplicate=True)

    async def create(request: Request):
        return await _save(request, K, None)

    async def bulk(request: Request):
        form = await request.form()
        ids = [int(x) for x in form.getlist("id") if x.isdigit()]
        try:
            n = fin(request).bulk_update(ids, form.get("status", ""), form.get("invoice_status", ""), form.get("day", ""))
        except ValueError as e:
            raise HTTPException(400, "Niepoprawna data.") from e
        log(request, f"{K['kind']}_bulk", f"{n}: {','.join(map(str, ids))[:200]}")
        return back(safe_path(form.get("wroc"), base), f"zmieniono-{n}")

    async def update(request: Request, oid: int):
        _get(request, K, oid)
        return await _save(request, K, oid)

    async def delete(request: Request, oid: int):
        form = await request.form()
        _get(request, K, oid)
        fin(request).delete_order(oid)
        log(request, f"{K['kind']}_deleted", str(oid))
        return back(safe_path(form.get("wroc"), base), "usunieto")

    router.add_api_route(base, list_, methods=["GET"], response_class=HTMLResponse)
    router.add_api_route(base + "/nowe", new, methods=["GET"], response_class=HTMLResponse)
    router.add_api_route(base + "/{oid:int}", edit, methods=["GET"], response_class=HTMLResponse)
    router.add_api_route(base + "/{oid:int}/duplikuj", duplicate, methods=["GET"], response_class=HTMLResponse)
    router.add_api_route(base, create, methods=["POST"])
    router.add_api_route(base + "/zbiorczo", bulk, methods=["POST"])
    router.add_api_route(base + "/{oid:int}", update, methods=["POST"])
    router.add_api_route(base + "/{oid:int}/usun", delete, methods=["POST"])


for _K in KINDS.values():
    register(_K)


def _parse_prospects(raw: str, F: Finance) -> list[dict]:
    try:
        items = json.loads(raw)
    except ValueError as e:
        raise ValueError("To nie jest poprawny JSON.") from e
    if isinstance(items, dict):
        items = items.get("prospects") or items.get("prospekty") or []
    if not isinstance(items, list):
        raise ValueError("Oczekuję listy prospektów.")
    out = []
    for i, it in enumerate(items):
        if not isinstance(it, dict) or not str(it.get("title", "")).strip() or not str(it.get("client", "")).strip():
            raise ValueError(f"Pozycja {i + 1}: brak tytułu albo klienta.")
        amount = it.get("amount_net")
        price = to_grosze(amount) if amount not in (None, "") else 0
        day = str(it.get("day") or date.today().isoformat())[:10]
        date.fromisoformat(day)
        client = F.find_client(str(it["client"]))
        note = " ".join(x for x in (str(it.get("note") or "").strip(), str(it.get("thread_url") or "").strip()) if x)
        out.append({"i": i, "title": str(it["title"]).strip()[:250], "client_name": str(it["client"]).strip()[:120],
                    "client": client, "principal": str(it.get("principal") or "").strip()[:120], "price": price,
                    "day": day, "note": note[:2000], "confidence": str(it.get("confidence") or "")})
    return out


@router.get("/prospekty/import", response_class=HTMLResponse)
async def prospects_import_form(request: Request):
    return page(request, "fin_prospekty_import.html", items=None, raw="", error="")


@router.post("/prospekty/import", response_class=HTMLResponse)
async def prospects_import(request: Request):
    form = await request.form()
    F = fin(request)
    raw = form.get("json", "")
    try:
        items = _parse_prospects(raw, F)
    except ValueError as e:
        return page(request, "fin_prospekty_import.html", items=None, raw=raw, error=str(e))
    if form.get("zapisz") != "1":
        return page(request, "fin_prospekty_import.html", items=items, raw=raw, error="")
    chosen = {int(x) for x in form.getlist("pick") if x.isdigit()}
    statuses = F.get_list("status")
    n = 0
    for it in items:
        if it["i"] not in chosen:
            continue
        cid = it["client"]["id"] if it["client"] else F.add_client(it["client_name"])
        F.save_order({"kind": "prospekt", "day": it["day"], "title": it["title"], "client_id": cid, "payer_id": cid,
                      "principal": it["principal"], "price": it["price"], "qty": 1,
                      "status": statuses[0] if statuses else "", "invoice_status": "Brak", "note": it["note"]})
        n += 1
    log(request, "prospects_imported", str(n))
    return back("/prospekty?wszystkie=1", f"zaimportowano-{n}")


@router.post("/prospekty/{oid:int}/do-zlecen")
async def prospect_accept(request: Request, oid: int):
    form = await request.form()
    _get(request, KINDS["prospekt"], oid)
    try:
        fin(request).accept_prospect(oid, form.get("day") or date.today().isoformat())
    except ValueError as e:
        raise HTTPException(400, "Niepoprawna data.") from e
    log(request, "prospect_accepted", str(oid))
    return RedirectResponse(f"/zlecenia/{oid}?ok=przeniesiono", status_code=303)


# ---------- klienci ----------
@router.get("/klienci", response_class=HTMLResponse)
async def clients(request: Request):
    return page(request, "fin_klienci.html", clients=fin(request).clients(), msg=request.query_params.get("ok", ""),
                error=request.query_params.get("blad", ""))


@router.post("/klienci")
async def client_add(request: Request):
    form = await request.form()
    try:
        fin(request).add_client(form.get("name", ""))
    except ValueError as e:
        return back(f"/klienci?blad={e}")
    return back("/klienci", "dodano")


@router.post("/klienci/{cid}")
async def client_rename(request: Request, cid: int):
    form = await request.form()
    try:
        fin(request).rename_client(cid, form.get("name", ""))
    except ValueError as e:
        return back(f"/klienci?blad={e}")
    return back("/klienci", "zapisano")


@router.post("/klienci/{cid}/usun")
async def client_delete(request: Request, cid: int):
    fin(request).delete_client(cid)
    log(request, "client_deleted", str(cid))
    return back("/klienci", "usunieto")


# ---------- koszty miesięczne ----------
@router.get("/koszty", response_class=HTMLResponse)
async def costs(request: Request, m: str = ""):
    month = month_or_now(m)
    F = fin(request)
    rows = F.costs(month)
    prev = prev_month(month)
    have = {(r["name"].strip().lower(), r["value"]) for r in rows}
    prev_rows = [{"r": r, "dup": (r["name"].strip().lower(), r["value"]) in have} for r in F.costs(prev)]
    return page(request, "fin_koszty.html", month=month, label=month_label(month), rows=rows,
                total=sum(r["value"] for r in rows), months=F.cost_months(), prev=prev, prev_rows=prev_rows,
                msg=request.query_params.get("ok", ""))


@router.post("/koszty")
async def cost_add(request: Request):
    form = await request.form()
    month = month_or_now(form.get("month", ""))
    try:
        fin(request).add_cost(month, form.get("name", ""), to_grosze(form.get("value")))
    except (ValueError, OverflowError) as e:
        raise HTTPException(400, str(e)) from e
    return back(f"/koszty?m={month}", "dodano")


@router.post("/koszty/kopiuj")
async def cost_copy(request: Request):
    form = await request.form()
    if not (valid_month(form.get("src", "")) and valid_month(form.get("dst", ""))):
        raise HTTPException(400, "Niepoprawny miesiąc.")
    ids = [int(x) for x in form.getlist("id") if x.isdigit()]
    if not ids:
        return back(f"/koszty?m={form.get('dst')}", "nic-nie-zaznaczono")
    n = fin(request).copy_costs(form.get("src", ""), form.get("dst", ""), ids)
    return back(f"/koszty?m={form.get('dst')}", f"skopiowano-{n}")


@router.post("/koszty/{cid}")
async def cost_update(request: Request, cid: int):
    form = await request.form()
    try:
        fin(request).update_cost(cid, form.get("name", ""), to_grosze(form.get("value")))
    except (ValueError, OverflowError) as e:
        raise HTTPException(400, str(e)) from e
    return back(f"/koszty?m={month_or_now(form.get('month', ''))}", "zapisano")


@router.post("/koszty/{cid}/usun")
async def cost_delete(request: Request, cid: int):
    form = await request.form()
    fin(request).delete_cost(cid)
    return back(f"/koszty?m={month_or_now(form.get('month', ''))}", "usunieto")


# ---------- podsumowania ----------
@router.get("/finanse", response_class=HTMLResponse)
async def summary(request: Request, m: str = ""):
    F = fin(request)
    month = month_or_now(m)
    return page(request, "fin_finanse.html", cur=F.month_summary(month), prev=F.month_summary(prev_month(month)),
                by_client=F.by_client(month), history=F.history(), years=F.year_table(), month=month,
                prev_m=prev_month(month))


# ---------- ustawienia statusów ----------
@router.get("/ustawienia", response_class=HTMLResponse)
async def settings_page(request: Request):
    F = fin(request)
    journal = request.app.state.db.q("SELECT * FROM audit ORDER BY id DESC LIMIT 40")
    return page(request, "fin_ustawienia.html", journal=journal, statuses="\n".join(F.get_list("status")),
                inv_statuses="\n".join(F.get_list("invoice_status")), msg=request.query_params.get("ok", ""))


@router.post("/ustawienia")
async def settings_save(request: Request):
    form = await request.form()
    F = fin(request)
    F.set_list("status", form.get("status", "").splitlines())
    F.set_list("invoice_status", form.get("invoice_status", "").splitlines())
    return back("/ustawienia", "zapisano")


# ---------- raporty dla klientów ----------
@router.get("/raporty", response_class=HTMLResponse)
async def reports(request: Request):
    F = fin(request)
    names = F.client_names()
    rows = [{"r": r, "f": Filters.from_query(json.loads(r["filters"]))} for r in F.reports()]
    return page(request, "fin_raporty.html", rows=rows, names=names, origin=request.app.state.cfg.origin,
                msg=request.query_params.get("ok", ""), new=request.query_params.get("nowy", ""))


@router.post("/raporty")
async def report_create(request: Request):
    form = await request.form()
    try:
        raw = json.loads(form.get("filters") or "{}")
        if not isinstance(raw, dict):
            raise ValueError("Niepoprawne filtry.")
        token = fin(request).create_report(Filters.from_query(raw), form.get("title", ""), form.get("expires") or None)
    except ValueError as e:
        msg = str(e) if "klienta" in str(e) else "Niepoprawne filtry albo data ważności."
        return RedirectResponse("/zlecenia?" + urlencode({**raw, "blad": msg}, doseq=True) if isinstance(raw, dict)
                                else "/zlecenia", status_code=303)
    log(request, "report_created", token[:8])
    return RedirectResponse(f"/raporty?nowy={token}", status_code=303)


@router.post("/raporty/{token}/wylacz")
async def report_revoke(request: Request, token: str):
    fin(request).revoke_report(token)
    log(request, "report_revoked", token[:8])
    return back("/raporty", "wylaczono")


def _report_data(request: Request, token: str):
    F = fin(request)
    r = F.report(token)
    if r is None:
        raise HTTPException(404, "Link wygasł albo jest nieaktywny.")
    f = Filters.from_query(json.loads(r["filters"]))
    return r, f, F.orders(f, limit=None), F.totals(f)


@router.get("/r/{token}", response_class=HTMLResponse)
async def report_public(request: Request, token: str):
    r, f, orders, totals = _report_data(request, token)
    return page(request, "fin_raport_publiczny.html", public=True, r=r, f=f, orders=orders, totals=totals, token=token)


@router.get("/r/{token}/xlsx")
async def report_xlsx(request: Request, token: str):
    r, f, orders, totals = _report_data(request, token)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Podsumowanie"
    ws.append(["Data", "Tytuł", "Zleceniodawca", "Status", "Cena netto", "Ilość", "Wartość netto"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for o in orders:
        ws.append([o["day"], o["title"], o["principal"], o["status"], o["price"] / 100, o["qty"], o["total"] / 100])
    ws.append([])
    ws.append(["", "Razem", "", "", "", "", totals["total"] / 100])
    ws[ws.max_row][6].font = Font(bold=True)
    for col, w in zip("ABCDEFG", (12, 50, 24, 16, 14, 8, 16)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=7):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = '# ##0.00'
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="Podsumowanie_brandwise.xlsx"'})


__all__ = ["router", "zl"]


# ---------- faktura ze zleceń (Fakturownia) ----------
def _invoice_ids(params) -> list[int]:
    return sorted({int(x) for x in params.getlist("id") if str(x).isdigit()})


def _invoice_check(orders) -> str:
    if not orders:
        return "Zaznacz zlecenia do zafakturowania."
    if any(o["kind"] != "zlecenie" for o in orders):
        return "Prospektów nie fakturujemy – najpierw dodaj je do zleceń."
    if len({o["payer_id"] for o in orders}) > 1:
        return "Zaznaczone zlecenia mają różnych płatników – jedna faktura może iść tylko do jednego płatnika."
    done = [o for o in orders if o["fk_invoice_id"] or o["invoice_status"] == "Faktura"]
    if done:
        return f"Część zleceń jest już zafakturowana ({', '.join(o['title'] for o in done[:3])}) – odznacz je."
    return ""


def _best_match(name: str, fk_clients: list[dict]) -> int | None:
    n = name.lower().strip()
    for c in fk_clients:
        if c["name"].lower().strip() == n:
            return c["id"]
    for c in fk_clients:
        if n and (n in c["name"].lower() or c["name"].lower() in n):
            return c["id"]
    return None


@router.get("/zlecenia/faktura", response_class=HTMLResponse)
async def invoice_preview(request: Request):
    F, cfg = fin(request), request.app.state.cfg
    orders = F.orders_by_ids(_invoice_ids(request.query_params))
    error = _invoice_check(orders)
    wroc = safe_path(request.query_params.get("wroc"))
    if error:
        return back(wroc + ("&" if "?" in wroc else "?") + urlencode({"blad": error}))
    fk = request.app.state.fakturownia
    try:
        fk_clients = await fk.clients()
    except Exception as e:  # noqa: BLE001
        return back(wroc + ("&" if "?" in wroc else "?") + urlencode({"blad": f"Fakturownia: {e}"}))
    payer = F.client(orders[0]["payer_id"])
    chosen = payer["fk_client_id"] or _best_match(payer["name"], fk_clients)
    net = sum(o["total"] for o in orders)
    vat = round(net * cfg.vat_rate / 100)
    return page(request, "fin_faktura.html", orders=orders, payer=payer, fk_clients=fk_clients, chosen=chosen,
                today=date.today().isoformat(), sell_date=max(o["day"] for o in orders),
                days=payer["payment_days"] or cfg.default_payment_days, vat_rate=cfg.vat_rate,
                net=net, vat=vat, gross=net + vat, wroc=wroc, remembered=bool(payer["fk_client_id"]))


@router.post("/zlecenia/faktura")
async def invoice_issue(request: Request):
    form = await request.form()
    F, cfg = fin(request), request.app.state.cfg
    wroc = safe_path(form.get("wroc"))
    orders = F.orders_by_ids(_invoice_ids(form))            # kwoty zawsze z bazy, nie z formularza
    error = _invoice_check(orders)
    try:
        fk_client_id = int(form.get("fk_client_id") or 0)
        days = int(form.get("days") or cfg.default_payment_days)
        issue = date.fromisoformat(form.get("issue_date", "")).isoformat()
        sell = date.fromisoformat(form.get("sell_date", "")).isoformat()
    except ValueError:
        error = error or "Niepoprawna data, termin albo nabywca."
    if not error and not fk_client_id:
        error = "Wybierz nabywcę z listy klientów Fakturowni."
    if error:
        return back(wroc + ("&" if "?" in wroc else "?") + urlencode({"blad": error}))
    positions = []
    for o in orders:
        name = (form.get(f"nazwa_{o['id']}") or o["title"]).strip()[:250]
        positions.append({"name": name, "quantity": f"{o['qty']:g}", "price_net": f"{o['price'] / 100:.2f}",
                          "tax": cfg.vat_rate})
    invoice = {"kind": "vat", "issue_date": issue, "sell_date": sell, "payment_to_kind": days,
               "client_id": fk_client_id, "positions": positions}
    fk = request.app.state.fakturownia
    try:
        res = await fk.create_invoice(invoice)
    except Exception as e:  # noqa: BLE001
        log(request, "invoice_failed", str(e)[:300])
        return back(wroc + ("&" if "?" in wroc else "?") + urlencode({"blad": str(e)[:400]}))
    inv_id, number = int(res.get("id") or 0) if isinstance(res, dict) else 0, str(res.get("number") or "") if isinstance(res, dict) else ""
    if not inv_id:                                           # bez id nie oznaczamy zleceń jako zafakturowane
        log(request, "invoice_failed", "brak id w odpowiedzi: " + json.dumps(res, ensure_ascii=False)[:300])
        return back(wroc + ("&" if "?" in wroc else "?") + urlencode(
            {"blad": "Fakturownia nie potwierdziła utworzenia faktury – sprawdź w Fakturowni, zanim spróbujesz ponownie."}))
    F.mark_invoiced([o["id"] for o in orders], inv_id, number)
    if form.get("zapamietaj"):
        F.set_client_billing(orders[0]["payer_id"], fk_client_id, days)
    expected = sum(o["total"] for o in orders)
    got = to_grosze(res.get("price_net")) if res.get("price_net") not in (None, "") else expected
    log(request, "invoice_issued", f"{number} id={inv_id} orders={','.join(str(o['id']) for o in orders)} net={got}")
    msg = f"Wystawiono fakturę {number or inv_id} na {zl(got)} netto."
    if got != expected:
        msg += f" UWAGA: Fakturownia policzyła netto {zl(got)}, a zlecenia dają {zl(expected)} – sprawdź fakturę."
    return RedirectResponse(f"/zlecenia/faktura/{inv_id}?" + urlencode({"msg": msg, "wroc": wroc}), status_code=303)


@router.get("/zlecenia/faktura/{inv_id}", response_class=HTMLResponse)
async def invoice_done(request: Request, inv_id: int):
    fk = request.app.state.fakturownia
    return page(request, "fin_faktura_ok.html", msg=request.query_params.get("msg", ""),
                url=f"{fk.base}/invoices/{inv_id}", wroc=safe_path(request.query_params.get("wroc")))
